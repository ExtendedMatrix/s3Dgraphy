"""TRAD1 — multilingual datamodel labels/descriptions (sidecar + xlsx round-trip).

The English source lives in the node datamodel and in the qualia vocabulary; a
sidecar `datamodel_translations.json` adds other languages with a per-key
`validated_<lang>` flag — node classes under `entries` (the shape EMStudio's
rules.ts reads), qualia / categories / subcategories in their own sections keyed
by id. This tests the mechanism: en is seeded from the sources and never lost,
the partners' xlsx round-trips translations and statuses, en is never
overwritten from the xlsx, and the reader falls back to English.
"""

import json
import shutil

import pytest

from s3dgraphy.tools import datamodel_i18n as t


@pytest.fixture
def sidecar(tmp_path, monkeypatch):
    """A temp copy of the real sidecar so the committed file is untouched."""
    tr = tmp_path / "datamodel_translations.json"
    shutil.copyfile(t.TRANSLATIONS, tr)
    monkeypatch.setattr(t, "TRANSLATIONS", tr)
    t._doc.cache_clear()
    yield tr
    t._doc.cache_clear()


def _xlsx_rows(path):
    from openpyxl import load_workbook
    wb = load_workbook(str(path))
    ws = wb[t.SHEET]
    hdr = [c.value for c in ws[1]]
    return wb, ws, {h: i for i, h in enumerate(hdr)}


def test_collect_en_covers_the_datamodel_classes():
    en = t._collect_en(t._load(t.DATAMODEL))
    assert len(en) >= 40
    su = en.get("StratigraphicUnit", {})
    assert su.get("description") and su.get("label")


def test_collect_covers_every_qualia_category_and_subcategory():
    q = t._load(t.QUALIA)
    en = t._collect_qualia_en(q)
    n_qualia = sum(len(s["qualia"]) for c in q["qualia_categories"]
                   for s in c["subcategories"].values())
    assert len(en["qualia"]) == n_qualia
    assert len(en["qualia_categories"]) == len(q["qualia_categories"])
    assert en["qualia"]["thickness"]["label"] == "Thickness"
    assert en["qualia_subcategories"]["dimensional"]["label"] == "Dimensional Qualia"


def test_committed_sidecar_en_is_in_sync():
    """The shipped sidecar's en must match the sources (── the `--check`)."""
    assert t.check() == 0


def test_committed_sidecar_keeps_the_emstudio_shape_and_all_languages():
    doc = t._load(t.TRANSLATIONS)
    assert doc["languages"] == t.LANGUAGES and "de" in doc["languages"]
    # EMStudio rules.ts: entries[class].description[locale]
    assert isinstance(doc["entries"]["StratigraphicUnit"]["description"]["en"], str)
    for section in t.SECTIONS:
        assert doc[section], section


def test_committed_sidecar_every_label_is_drafted_in_every_language():
    """Decision E.D. 27 Sep: all languages written now, unverified until the
    partners confirm. Labels (nodes + qualia) and node descriptions."""
    doc = t._load(t.TRANSLATIONS)
    missing = [(s, k, f, lang) for s, k, f in t._rows(doc)
               for lang in t.LANGUAGES
               if not str(doc[s][k][f].get(lang) or "").strip()]
    assert missing == []
    cov = t.coverage(doc)
    assert all(c["present"] == c["total"] for c in cov.values())


def test_round_trip_status_and_en_not_overwritten(sidecar, tmp_path):
    doc = json.loads(sidecar.read_text())
    source_en = doc["qualia"]["thickness"]["label"]["en"]
    xlsx = tmp_path / "t.xlsx"
    t.export_xlsx(str(xlsx))

    wb, ws, ci = _xlsx_rows(xlsx)
    for row in ws.iter_rows(min_row=2):
        if row[ci["key"]].value == "qualia.thickness.label":
            row[ci["en (source)"]].value = "HACKED EN"          # must be ignored
            row[ci["it"]].value = "Spessore (corretto)"
            row[ci["status it"]].value = "verified"
            row[ci["status de"]].value = "needs change"
            row[ci["reviewer"]].value = "CNR-ISPC"
            row[ci["comment"]].value = "de: Mächtigkeit?"
    wb.save(str(xlsx))

    t.import_xlsx(str(xlsx))

    out = json.loads(sidecar.read_text())["qualia"]["thickness"]["label"]
    assert out["en"] == source_en                  # en never overwritten from xlsx
    assert out["it"] == "Spessore (corretto)"      # translation round-tripped
    assert out["validated_it"] is True             # verified → validated_it: true
    assert out["validated_de"] is False and out["status_de"] == "needs change"
    assert out["reviewer"] == "CNR-ISPC" and out["comment"] == "de: Mächtigkeit?"

    # …and back out again: the status survives a second export
    xlsx2 = tmp_path / "t2.xlsx"
    t.export_xlsx(str(xlsx2))
    _, ws2, ci2 = _xlsx_rows(xlsx2)
    row = next(r for r in ws2.iter_rows(min_row=2, values_only=True)
               if r[ci2["key"]] == "qualia.thickness.label")
    assert row[ci2["status it"]] == "verified"
    assert row[ci2["status de"]] == "needs change"
    assert row[ci2["status el"]] == "draft"


def test_export_refuses_to_overwrite_and_sync_imports_first(sidecar, tmp_path):
    xlsx = tmp_path / "partners.xlsx"
    t.export_xlsx(str(xlsx))
    with pytest.raises(FileExistsError):
        t.export_xlsx(str(xlsx))

    wb, ws, ci = _xlsx_rows(xlsx)
    for row in ws.iter_rows(min_row=2):
        if row[ci["key"]].value == "node.EpochNode.label":
            row[ci["status pl"]].value = "verified"
    wb.save(str(xlsx))

    t.sync_xlsx(str(xlsx))  # import the partners' work, back up, re-export

    assert json.loads(sidecar.read_text())["entries"]["EpochNode"]["label"]["validated_pl"] is True
    _, ws2, ci2 = _xlsx_rows(xlsx)
    row = next(r for r in ws2.iter_rows(min_row=2, values_only=True)
               if r[ci2["key"]] == "node.EpochNode.label")
    assert row[ci2["status pl"]] == "verified"
    assert list(tmp_path.glob("partners_backup_*.xlsx"))


def test_seed_is_idempotent_and_preserves_translations(sidecar):
    doc = json.loads(sidecar.read_text())
    doc["entries"]["StratigraphicUnit"]["description"]["it"] = "Preservami."
    doc["qualia"]["thickness"]["label"]["de"] = "Preservami."
    doc["qualia"]["thickness"]["label"]["validated_de"] = True
    sidecar.write_text(json.dumps(doc, ensure_ascii=False))

    t.seed()  # re-seed en

    out = json.loads(sidecar.read_text())
    assert out["entries"]["StratigraphicUnit"]["description"]["it"] == "Preservami."
    assert out["qualia"]["thickness"]["label"]["de"] == "Preservami."
    assert out["qualia"]["thickness"]["label"]["validated_de"] is True
    assert t.check() == 0


def test_reader_falls_back_to_english(sidecar):
    doc = json.loads(sidecar.read_text())
    doc["qualia"]["thickness"]["label"].pop("pl", None)
    sidecar.write_text(json.dumps(doc, ensure_ascii=False))
    t._doc.cache_clear()

    assert t.qualia_label("thickness", "pl") == "Thickness"      # missing text
    assert t.qualia_label("thickness", "xx") == "Thickness"      # unknown language
    assert t.qualia_label("no_such_quale", "it") is None
    assert t.qualia_label("thickness", "it-IT") == t.qualia_label("thickness", "it")


def test_reader_on_the_committed_sidecar():
    t._doc.cache_clear()
    assert t.qualia_label("thickness", "it") == "Spessore"       # ICCD lexicon
    for lang in ("he", "de"):
        lab = t.qualia_label("thickness", lang)
        assert lab and lab != "Thickness"
        assert t.is_validated("qualia", "thickness", "label", lang) is False  # a draft
    assert t.qualia_category_label("physical_material", "en") == "Physical and Material Qualia"
    assert t.node_label("EpochNode", "it")


def test_property_names_are_labelled_and_are_not_qualia():
    """em_qualia_types 1.6.8 / translations 1.9 (MICRO badge PD e cronologia,
    part 5): the EM property names in use that are not a quale — `definition`,
    `material` — have their own section, seeded from the qualia file's
    `property_names` block, labelled in every language, it/de validated; a
    name of that block is never a quale id, and `_` keys are notes."""
    t._doc.cache_clear()
    qualia = json.loads(t.QUALIA.read_text(encoding="utf-8"))
    en = t._collect_property_names_en(qualia)
    assert set(en) == {"definition", "material"}
    assert not any(k.startswith("_") for k in en)
    ids = set(t._collect_qualia_en(qualia)["qualia"])
    assert not (set(en) & ids)
    assert qualia["property_names"]["material"]["quale"] in ids
    assert t.property_name_label("material", "it") == "Materiale"
    assert t.property_name_label("definition", "it") == "Definizione"
    assert t.property_name_label("material", "xx") == "Material"
    assert t.property_name_label("riser_depth", "it") is None
    for name in en:
        for lang in ("it", "de"):
            assert t.is_validated("property_names", name, "label", lang) is True
    # E.D.'s words for the dating qualia
    assert t.qualia_label("absolute_time_start", "it") == "Data di inizio"
    assert t.qualia_label("absolute_time_end", "it") == "Data di fine"
