"""Human authorship with AI support, verified on NODES (connections 1.6.25).

E.D. 2026-09-29: an AI proposal a person accepts stays signed by the person,
with the mark «supporto AI»; it stays among the warnings until a person
verifies it with their ORCID. Same names as the narrative blocks.
"""

import pytest

from s3dgraphy import api
from s3dgraphy.ai_validation import (
    AIValidationError, UNVALIDATED_MARK, ai_marker, export_view,
    is_unvalidated_ai, mark_ai_assisted, unvalidated_ai, validate_node,
)
from s3dgraphy.crdt import META_KEYS
from s3dgraphy.graph import Graph
from s3dgraphy.nodes import StratigraphicUnit, PropertyNode
from s3dgraphy.nodes.author_node import AuthorAINode, AuthorNode

ORCID = "0000-0002-1825-0097"


def _graph() -> Graph:
    g = Graph(graph_id="ai")
    g.add_node(AuthorNode("ed", name="Emanuel", orcid=ORCID, surname="D"))
    g.add_node(AuthorNode("anon", name="Anon"))              # no ORCID
    g.add_node(AuthorAINode("claude", name="Claude"))
    g.add_node(StratigraphicUnit("US5", name="US5", description="muro"))
    g.add_node(StratigraphicUnit("US6", name="US6", description="limo sabbioso"))
    g.add_node(PropertyNode("p1", name="material", value="tufo",
                            property_type="material"))
    g.add_edge("h1", "US5", "p1", "has_property")
    mark_ai_assisted(g, g.find_node_by_id("p1"), by="claude",
                     model="claude-opus-5-5", prompt_ref=None)
    mark_ai_assisted(g, g.find_node_by_id("US6"), by="claude",
                     fields=["description"])
    return g


# ── the marker and the verification ─────────────────────────────────────────

def test_marker_shape_and_unvalidated_list():
    g = _graph()
    assert g.find_node_by_id("p1").data["ai_assisted"] == {
        "by": "claude", "model": "claude-opus-5-5"}
    rows = {r["node"]: r for r in unvalidated_ai(g)}
    assert set(rows) == {"p1", "US6"}
    assert rows["p1"]["fields"] is None                 # the whole node
    assert rows["US6"]["fields"] == ["description"]     # one field
    assert api.unvalidated_ai(g) == unvalidated_ai(g)


def test_marker_needs_an_ai_author():
    g = _graph()
    with pytest.raises(AIValidationError):
        mark_ai_assisted(g, g.find_node_by_id("US5"), by="ed")


def test_only_a_person_with_orcid_validates():
    g = _graph()
    with pytest.raises(AIValidationError):
        validate_node(g, g.find_node_by_id("p1"), "claude")   # a model
    with pytest.raises(AIValidationError):
        validate_node(g, g.find_node_by_id("p1"), "anon")     # no ORCID
    with pytest.raises(AIValidationError):
        validate_node(g, g.find_node_by_id("US5"), "ed")      # not AI-made
    out = api.validate_ai(g, "p1", "ed", at="2026-10-04T10:00:00Z")
    assert out == {"validated_by": "ed", "validated_at": "2026-10-04T10:00:00Z"}
    assert not is_unvalidated_ai(g.find_node_by_id("p1"))
    assert [r["node"] for r in unvalidated_ai(g)] == ["US6"]


def test_marking_again_clears_the_verification():
    g = _graph()
    api.validate_ai(g, "p1", "ed")
    mark_ai_assisted(g, g.find_node_by_id("p1"), by="claude")
    assert is_unvalidated_ai(g.find_node_by_id("p1"))
    assert "validated_at" not in g.find_node_by_id("p1").data


def test_the_block_word_is_an_alias():
    n = StratigraphicUnit("u", name="u")
    n.data = {}
    n.data.update({"ai_generated": True, "authored_by": "claude",
                   "prompt_ref": "pr"})
    assert ai_marker(n) == {"by": "claude", "prompt_ref": "pr"}
    assert is_unvalidated_ai(n)


def test_block_carries_validated_at_like_a_node():
    from s3dgraphy.nodes.narrative_node import Block
    b = Block.ai_prose("x", author_id="claude")
    b.endorse("ed", at="2026-10-04T10:00:00Z")
    d = b.to_dict()
    assert d["validated_by"] == "ed" and d["validated_at"] == "2026-10-04T10:00:00Z"
    assert Block.from_dict(d).validated_at == "2026-10-04T10:00:00Z"


# ── the CRDT merges them as content ────────────────────────────────────────

def test_verification_fields_are_content_for_the_crdt():
    for key in ("ai_assisted", "validated_by", "validated_at"):
        assert key not in META_KEYS
    section = {"nodes": [{"id": "p1", "node_type": "property", "name": "m",
                          "data": {"ai_assisted": {"by": "claude"}}}],
               "edges": []}
    r = api.apply_op(section, api.make_op(
        "update_field", node_id="p1", field="data.validated_by", value="ed",
        ts="2026-10-04T10:00:00Z", author=ORCID))
    assert r["applied"] is True
    assert section["nodes"][0]["data"]["validated_by"] == "ed"
    # an older op does not undo a newer verification
    api.apply_op(section, api.make_op(
        "update_field", node_id="p1", field="data.validated_by", value="x",
        ts="2026-10-03T10:00:00Z", author=ORCID))
    assert section["nodes"][0]["data"]["validated_by"] == "ed"


# ── the export view ────────────────────────────────────────────────────────

def test_view_default_drops_whole_nodes_and_empties_fields():
    g = _graph()
    view, rows = export_view(g)
    ids = {n.node_id for n in view.nodes}
    assert "p1" not in ids and "US6" in ids
    assert not any(e.edge_target == "p1" for e in view.edges)
    assert view.find_node_by_id("US6").description == ""
    # the graph itself is untouched
    assert g.find_node_by_id("US6").description == "limo sabbioso"
    assert g.find_node_by_id("p1") is not None


def test_view_forced_marks_the_fields():
    view, rows = export_view(_graph(), include_unvalidated=True)
    assert view.find_node_by_id("p1").value == f"{UNVALIDATED_MARK} tufo"
    assert view.find_node_by_id("US6").description == f"{UNVALIDATED_MARK} limo sabbioso"
    assert view.find_node_by_id("US6").name == "US6"
    assert len(rows) == 2


def test_nothing_to_withhold_is_the_same_graph():
    g = Graph(graph_id="clean")
    g.add_node(StratigraphicUnit("US1", name="US1"))
    assert export_view(g) == (g, [])


# ── xlsx ───────────────────────────────────────────────────────────────────

def _claims(path):
    import openpyxl
    return [r for r in openpyxl.load_workbook(path)["Claims"].iter_rows(
        values_only=True)][1:]


def test_xlsx_default_leaves_out_and_forced_marks(tmp_path):
    pytest.importorskip("openpyxl")
    from s3dgraphy.exporter.unified_xlsx_exporter import UnifiedXLSXExporter
    from s3dgraphy.importer.unified_xlsx_importer import UnifiedXLSXImporter
    import openpyxl

    ex = UnifiedXLSXExporter(_graph())
    ex.write(str(tmp_path / "d.xlsx"))
    assert {r["node"] for r in ex.excluded} == {"p1", "US6"}
    assert not any(r[2] == "material" for r in _claims(tmp_path / "d.xlsx"))

    UnifiedXLSXExporter(_graph(), include_unvalidated=True).write(
        str(tmp_path / "f.xlsx"))
    rows = [r for r in _claims(tmp_path / "f.xlsx") if r[2] == "material"]
    assert rows and rows[0][3] == f"{UNVALIDATED_MARK} tufo"
    units = {r[0]: r for r in openpyxl.load_workbook(tmp_path / "f.xlsx")[
        "Units"].iter_rows(values_only=True)}
    assert units["US6"][2] == f"{UNVALIDATED_MARK} limo sabbioso"

    # read back: the mark comes off the value and stays on the node
    g2 = UnifiedXLSXImporter(str(tmp_path / "f.xlsx"), graph_id="ai").parse()
    pn = next(n for n in g2.nodes if isinstance(n, PropertyNode))
    assert pn.value == "tufo" and is_unvalidated_ai(pn)
    us6 = next(n for n in g2.nodes if n.name == "US6")
    assert us6.description == "limo sabbioso" and is_unvalidated_ai(us6)


# ── RDF ────────────────────────────────────────────────────────────────────

def test_rdf_default_leaves_out(tmp_path):
    pytest.importorskip("rdflib")
    from s3dgraphy.exporter.rdf_exporter import RDFExporter
    ex = RDFExporter(str(tmp_path / "d.ttl"), format="turtle")
    ttl = open(ex.export_single_graph(_graph()), encoding="utf-8").read()
    assert "tufo" not in ttl and "limo sabbioso" not in ttl
    assert ex.stats["ai_unvalidated"] == 2


def test_rdf_forced_roundtrip_keeps_the_marker_and_is_isomorphic(tmp_path):
    rdflib = pytest.importorskip("rdflib")
    from rdflib.compare import isomorphic
    from s3dgraphy.exporter.rdf_exporter import RDFExporter
    from s3dgraphy.importer.rdf_importer import RDFImporter

    g = _graph()
    api.validate_ai(g, "p1", "ed", at="2026-10-04T10:00:00Z")
    t1 = RDFExporter(str(tmp_path / "a.ttl"), format="turtle",
                     include_unvalidated=True).export_single_graph(g)
    text = open(t1, encoding="utf-8").read()
    assert f"{UNVALIDATED_MARK} limo sabbioso" in text      # unvalidated: marked
    assert f"{UNVALIDATED_MARK} tufo" not in text    # validated: clean
    g2 = RDFImporter().parse(t1)[0]
    p1 = g2.find_node_by_id("p1")
    assert p1.data["ai_assisted"] == {"by": "claude", "model": "claude-opus-5-5"}
    assert p1.data["validated_by"] == "ed"
    assert p1.data["validated_at"] == "2026-10-04T10:00:00Z"
    us6 = g2.find_node_by_id("US6")
    assert us6.description == "limo sabbioso"
    assert us6.data["ai_assisted"] == {"by": "claude", "fields": ["description"]}
    # the marker triples are not read back as edges
    assert not any(e.edge_target in ("claude", "ed") and e.edge_source in ("p1", "US6")
                   for e in g2.edges)
    t2 = RDFExporter(str(tmp_path / "b.ttl"), format="turtle",
                     include_unvalidated=True).export_single_graph(g2)
    a, b = rdflib.Graph(), rdflib.Graph()
    a.parse(t1, format="turtle"); b.parse(t2, format="turtle")
    assert isomorphic(a, b)


# ── html (through the bake) ─────────────────────────────────────────────────

def _with_narrative(g):
    from s3dgraphy.nodes.narrative_node import NarrativeNode, Chapter, Block
    n = NarrativeNode("nar", name="Racconto")
    n.chapters = [Chapter(title="Uno", blocks=[
        Block.prose("Il muro [[US5]] e lo strato [[US6]]."),
        Block.embed("US6", "us"),
        Block.embed("p1", "us"),
    ])]
    g.add_node(n)
    return g


def test_html_default_leaves_out_and_forced_marks():
    g = _with_narrative(_graph())
    baked = api.bake_narrative(g, "nar")
    assert {r["node"] for r in baked.excluded_nodes} == {"p1", "US6"}
    html = api.export_narrative_html(g, "nar")
    assert "tufo" not in html and "limo sabbioso" not in html
    forced = api.export_narrative_html(g, "nar", include_unvalidated=True)
    assert f"{UNVALIDATED_MARK} limo sabbioso" in forced
    assert "non validati da una persona" in forced
