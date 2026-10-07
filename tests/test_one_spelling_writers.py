"""Una grafia sola: whoever WRITES a physical bond writes the canonical name;
whoever READS one asks the datamodel.

The connections datamodel (1.6.20) declares `is_bonded_to` and
`is_physically_equal_to` as `spelling_of` `bonded_to` and `equals`: the same
relation, the same RDF projection, accepted when read, never written. Before
2026-09-27 five places in the code each kept their own list and knew one
spelling only (report 2026-10-18-MICRO-i-nomi-dei-legami-fisici, §6), so a
graph from pyArchInit (older spelling) and one from xlsx (canonical) disagreed
in silence. Each of the five is exercised here with BOTH spellings.

The pairs are not typed here: they are read off the datamodel, so the day a
third spelling is declared these tests cover it without being edited.
"""

from __future__ import annotations

import os
import re

import pytest
import rdflib

from s3dgraphy import api
from s3dgraphy.edges.connections_loader import get_connections_datamodel
from s3dgraphy.graph import Graph
from s3dgraphy.nodes.property_node import PropertyNode
from s3dgraphy.nodes.stratigraphic_node import StratigraphicUnit


def _pairs():
    """`[(older_spelling, canonical)]`, from the datamodel's `spelling_of`."""
    dm = get_connections_datamodel()
    return sorted((name, entry["spelling_of"])
                  for name, entry in dm._canonical_edges.items()
                  if entry.get("spelling_of"))


PAIRS = _pairs()
#: every spelling, labelled with the canonical it must come out as
SPELLINGS = [(old, canon) for old, canon in PAIRS] + \
            [(canon, canon) for _old, canon in PAIRS]


def test_the_datamodel_declares_the_two_pairs():
    assert PAIRS == [("is_bonded_to", "bonded_to"),
                     ("is_physically_equal_to", "equals")]


def _units(graph, *names):
    for name in names:
        graph.add_node(StratigraphicUnit(name, name=name))


# ── C1 · unified_xlsx_exporter ───────────────────────────────────────────────

def _claims(path):
    from openpyxl import load_workbook
    ws = load_workbook(path)["Claims"]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, [c.value for c in r])) for r in ws.iter_rows(min_row=2)]
    return [(r["TARGET_ID"], r["PROPERTY_TYPE"], r["TARGET2_ID"]) for r in rows
            if r["PROPERTY_TYPE"] not in (None, "has_first_epoch",
                                          "belongs_to_epoch")]


@pytest.mark.parametrize("spelling,canonical", SPELLINGS)
def test_the_xlsx_exporter_writes_the_canonical_name(tmp_path, spelling, canonical):
    """Until 2026-09-27 an older spelling left 0 claims, and no warning."""
    from s3dgraphy.exporter.unified_xlsx_exporter import UnifiedXLSXExporter
    g = Graph(graph_id="x")
    _units(g, "US1", "US2")
    g.add_edge("e", "US1", "US2", spelling)
    out = tmp_path / "em_data.xlsx"
    UnifiedXLSXExporter(g).write(str(out))
    assert _claims(str(out)) == [("US1", canonical, "US2")]


def test_one_bond_under_two_spellings_is_one_claim(tmp_path):
    """The case the double arc made: the same fact, once per spelling and from
    both ends of a directionless relation — one row, not three."""
    from s3dgraphy.exporter.unified_xlsx_exporter import UnifiedXLSXExporter
    g = Graph(graph_id="x")
    _units(g, "US1", "US2")
    g.add_edge("a", "US1", "US2", "is_bonded_to")
    g.add_edge("b", "US1", "US2", "bonded_to")
    g.add_edge("c", "US2", "US1", "bonded_to")
    out = tmp_path / "em_data.xlsx"
    UnifiedXLSXExporter(g).write(str(out))
    assert len(_claims(str(out))) == 1, _claims(str(out))


def test_a_pyarchinit_graph_keeps_every_bond_on_the_way_to_xlsx(tmp_path):
    """The grafo pyArchInit → xlsx: `rapporti` in all ten UI languages, parsed
    into edges, exported, re-imported — every «Si lega a» and «Uguale a» is
    still there, as the canonical edge. Plus a stored graph from BEFORE the
    repair, whose edges still carry the older spelling."""
    from s3dgraphy.exporter.unified_xlsx_exporter import UnifiedXLSXExporter
    from s3dgraphy.sync import rapporti

    g = Graph(graph_id="pyarchinit")
    expected = set()
    n = 0
    for lang, terms in sorted(rapporti._REL_TERMS_BY_LANG.items()):
        for term in terms[:2]:                     # 0 Uguale a · 1 Si lega a
            src, tgt = f"US{n}a", f"US{n}b"
            n += 1
            _units(g, src, tgt)
            (edge_type, *_rest), = rapporti.parse_rapporti([[term, tgt]])
            g.add_edge(f"r{n}", src, tgt, edge_type)
            expected.add((src, edge_type, tgt))
    for old, canon in PAIRS:                       # a graph stored before
        src, tgt = f"OLD_{old}_a", f"OLD_{old}_b"
        _units(g, src, tgt)
        g.add_edge(f"old_{old}", src, tgt, old)
        expected.add((src, canon, tgt))

    out = tmp_path / "em_data.xlsx"
    UnifiedXLSXExporter(g).write(str(out))
    assert set(_claims(str(out))) == expected

    again, _warnings, _stats = api.em_data_to_graph(str(out), graph_id="again")
    names = {nd.node_id: nd.name for nd in again.nodes}
    bonds = {(names[e.edge_source], e.edge_type, names[e.edge_target])
             for e in again.edges if e.edge_type in {c for _o, c in PAIRS}}
    assert bonds == expected


# ── C2 · unified_xlsx_importer ───────────────────────────────────────────────

def _workbook(path, property_type):
    from openpyxl import Workbook
    wb = Workbook()
    wb.remove(wb.active)
    for name, header, rows in (
        ("Units", ("ID", "TYPE", "NAME"),
         [("US1", "US", "one"), ("US2", "US", "two")]),
        ("Epochs", ("ID", "NAME", "START", "END", "COLOR"), []),
        ("Claims", ("TARGET_ID", "PROPERTY_TYPE", "VALUE", "TARGET2_ID"),
         [("US1", property_type, "", "US2")]),
        ("Authors", ("ID", "KIND", "DISPLAY_NAME", "ORCID", "AFFILIATION"), []),
        ("Documents", ("ID", "FILENAME", "TITLE", "YEAR", "AUTHOR_IDS",
                       "ROLE", "CONTENT_NATURE", "GEOMETRY"), []),
    ):
        ws = wb.create_sheet(name)
        ws.append(header)
        for row in rows:
            ws.append(row)
    wb.save(str(path))
    return str(path)


@pytest.mark.parametrize("spelling,canonical", SPELLINGS)
def test_the_xlsx_importer_makes_the_canonical_edge(tmp_path, spelling, canonical):
    """Until 2026-09-27 an older spelling became a PropertyNode named after it."""
    graph, _w, _s = api.em_data_to_graph(
        _workbook(tmp_path / "em_data.xlsx", spelling), graph_id="x")
    bonds = [e.edge_type for e in graph.edges
             if e.edge_type in get_connections_datamodel().spellings(canonical)]
    assert bonds == [canonical]
    assert not [nd for nd in graph.nodes if isinstance(nd, PropertyNode)
                and spelling in (nd.name, nd.property_type)]


# ── C3 · graph_merger ────────────────────────────────────────────────────────

def _merge_conflicts(existing_type, incoming_type):
    from s3dgraphy.merge.graph_merger import GraphMerger
    existing, incoming = Graph(graph_id="a"), Graph(graph_id="b")
    for g, et in ((existing, existing_type), (incoming, incoming_type)):
        _units(g, "US1", "US2")
        g.add_edge("e", "US1", "US2", et)
    return [c for c in GraphMerger().compare(existing, incoming)
            if c.field.startswith("edge")]


@pytest.mark.parametrize("old,canonical", PAIRS)
def test_two_spellings_of_one_bond_are_not_a_conflict(old, canonical):
    """Measured 2026-10-18: `edge_added` + `edge_removed` for one fact."""
    assert _merge_conflicts(old, canonical) == []
    assert _merge_conflicts(canonical, old) == []
    assert _merge_conflicts(canonical, canonical) == []
    assert _merge_conflicts(old, old) == []


def test_a_real_change_of_relation_is_still_a_conflict():
    """The control: normalising spellings must not blind the merger."""
    kinds = sorted(c.conflict_type for c in
                   _merge_conflicts("is_bonded_to", "overlies"))
    assert kinds == ["edge_added", "edge_removed"]


def test_removing_a_bond_finds_it_under_either_spelling():
    """The apply step looks the edge up by every spelling the datamodel
    accepts: the conflict names `bonded_to`, the graph holds `is_bonded_to`."""
    from s3dgraphy.merge.graph_merger import GraphMerger
    existing, incoming = Graph(graph_id="a"), Graph(graph_id="b")
    _units(existing, "US1", "US2")
    _units(incoming, "US1", "US2")
    existing.add_edge("e", "US1", "US2", "is_bonded_to")
    merger = GraphMerger()
    conflicts = merger.compare(existing, incoming)
    assert [c.field for c in conflicts if c.conflict_type == "edge_removed"] \
        == ["edge:bonded_to"]
    for c in conflicts:
        c.resolved = c.accepted = True
    merger.apply_resolutions(existing, conflicts, incoming)
    assert not [e for e in existing.edges if "bonded" in e.edge_type]


# ── C4 · inference_engine + graphml_exporter ─────────────────────────────────

@pytest.mark.parametrize("spelling,canonical", SPELLINGS)
def test_the_inference_engine_sees_either_spelling(spelling, canonical):
    """Measured 2026-10-18: 1 ambiguous relation for `is_bonded_to`, 0 for
    `bonded_to`."""
    from s3dgraphy.temporal.inference_engine import TemporalInferenceEngine
    g = Graph(graph_id="x")
    _units(g, "US1", "US2")
    g.add_edge("e", "US1", "US2", spelling)
    ambiguous = TemporalInferenceEngine().get_ambiguous_relations(g)
    assert [e.edge_type for e in ambiguous] == [spelling]


@pytest.mark.parametrize("spelling,canonical", SPELLINGS)
def test_graphml_turns_either_spelling_into_has_same_time(
        tmp_path, caplog, spelling, canonical):
    """`bonded_to` used to leave raw (in the yEd body the difference cannot be
    seen: same line), so the count the exporter reports is what is read — in
    its log since #27, no longer on stdout."""
    import logging
    from s3dgraphy.exporter.graphml.graphml_exporter import GraphMLExporter
    g = Graph(graph_id="x")
    _units(g, "US1", "US2", "US3")
    g.add_edge("e1", "US1", "US2", spelling)
    g.add_edge("e2", "US1", "US3", "overlies")
    with caplog.at_level(logging.DEBUG, logger="s3dgraphy.exporter.graphml"):
        GraphMLExporter(g).export(str(tmp_path / "x.graphml"))
    line = next(ln for ln in caplog.messages
                if ln.startswith("Generating") and "has_same_time" in ln)
    assert re.search(r"Generating 2 edges \(1 is_after, 1 has_same_time", line), line


@pytest.mark.parametrize("spelling,canonical", SPELLINGS)
def test_the_relations_string_names_the_canonical(spelling, canonical):
    from s3dgraphy.exporter.graphml.graphml_exporter import GraphMLExporter
    g = Graph(graph_id="x")
    _units(g, "US1", "US2")
    g.add_edge("e", "US1", "US2", spelling)
    exporter = GraphMLExporter(g)
    assert exporter._compute_relations_string(
        g.find_node_by_id("US1")) == f"{canonical}: US2"


# ── C5 · sync/rapporti ───────────────────────────────────────────────────────

def test_rapporti_parse_writes_the_canonical_in_every_language():
    from s3dgraphy.sync import rapporti
    canonical = {c for _o, c in PAIRS}
    older = {o for o, _c in PAIRS}
    for lang, terms in rapporti._REL_TERMS_BY_LANG.items():
        written = {rapporti.parse_rapporti([[t, "2"]])[0][0] for t in terms}
        assert canonical <= written and not written & older, lang


@pytest.mark.parametrize("spelling,canonical", SPELLINGS)
def test_rapporti_serialize_renders_either_spelling(spelling, canonical):
    """Measured 2026-10-18: `bonded_to` landed in d13 as
    `['bonded_to', 'B', …]` instead of `['Si lega a', …]`."""
    from s3dgraphy.sync import rapporti
    label = {"bonded_to": "Si lega a", "equals": "Uguale a"}[canonical]
    g = Graph(graph_id="x")
    _units(g, "US1", "US2")
    g.add_edge("e", "US1", "US2", spelling)
    out = rapporti.serialize_rapporti_from_edges(g, "Sito")
    # the target is written bare, prefix stripped — pyArchInit's own form
    assert out == {"US1": [[label, "2", "1", "Sito"]]}
    assert rapporti.select_rapporti_label(spelling, "US", "US") == label


def test_rapporti_round_trips_its_own_label():
    """parse → edge → serialise gives back the term the row had, in the
    row's own language (the node's stored `rapporti` anchors it)."""
    from s3dgraphy.sync import rapporti
    g = Graph(graph_id="x")
    _units(g, "US1", "US2")
    g.find_node_by_id("US1").attributes["rapporti"] = [["Bindet an", "2"]]
    (edge_type, *_r), = rapporti.parse_rapporti([["Bindet an", "2"]])
    assert edge_type == "bonded_to"
    g.add_edge("e", "US1", "US2", edge_type)
    assert rapporti.serialize_rapporti_from_edges(g, "S")["US1"][0][0] == \
        "Bindet an"


# ── the RDF projection did not move ─────────────────────────────────────────

def test_the_rdf_projection_is_the_same_for_both_spellings(tmp_path):
    """One triple per bond, whatever spelling the graph holds: 2 `em:bondedTo`
    and 2 `em:physicallyEquals` for a graph with both spellings of both.
    Measured against the tree before this change: the N-Triples are
    identical, 45 triples (report 2026-10-19)."""
    from s3dgraphy.exporter.rdf_exporter import RDFExporter
    g = Graph(graph_id="both")
    _units(g, *"ABCDEFGH")
    g.add_edge("1", "A", "B", "is_bonded_to")
    g.add_edge("2", "C", "D", "bonded_to")
    g.add_edge("3", "E", "F", "is_physically_equal_to")
    g.add_edge("4", "G", "H", "equals")
    path = str(tmp_path / "x.ttl")
    RDFExporter(path, format="turtle").export_single_graph(g)
    r = rdflib.Graph()
    r.parse(path, format="turtle")
    em = "https://w3id.org/em/ontology#"
    for local in ("bondedTo", "physicallyEquals"):
        assert len(list(r.triples((None, rdflib.URIRef(em + local), None)))) == 2
