"""A property with more than one owner (connections 1.6.23, ownership.py).

E.D., revisione della scrivania v7: a USV/s may instantiate the ``material`` of
the US it completes. One property, several owners, the inheritance DECLARED on
the edge; the property stays in the original owner's paradata group. These
tests pin that no projection loses an owner, and that the original owner is
still readable after the round trip.
"""

import pytest

from s3dgraphy import api
from s3dgraphy.graph import Graph
from s3dgraphy.nodes import (
    StratigraphicUnit, PropertyNode, DocumentNode, ExtractorNode,
)
from s3dgraphy.nodes.group_node import ParadataNodeGroup
from s3dgraphy.ownership import (
    inherit_property, original_owner, property_owners, shared_properties,
)
from s3dgraphy.utils.utils import get_stratigraphic_node_class

USVs = get_stratigraphic_node_class("USVs")


def _graph() -> Graph:
    g = Graph(graph_id="owners")
    g.add_node(StratigraphicUnit("US5", name="US5"))
    g.add_node(USVs("USV1", name="USV1"))
    g.add_node(PropertyNode("p1", name="material", value="tufo",
                            property_type="material"))
    g.add_node(DocumentNode("D1", name="D.1"))
    g.add_node(ExtractorNode("ext1", name="D.1.1", source="D.1 p.3"))
    g.add_edge("e1", "US5", "p1", "has_property")
    g.add_edge("e3", "p1", "ext1", "has_data_provenance")
    g.add_edge("e4", "ext1", "D1", "extracted_from")
    inherit_property(g, "USV1", "p1")
    return g


def _owners(g, name="material"):
    pns = [n for n in g.nodes if isinstance(n, PropertyNode) and n.name == name]
    assert len(pns) == 1, f"the property was split into {len(pns)}"
    return [(g.find_node_by_id(r["owner"]).name, r["original"])
            for r in property_owners(g, pns[0].node_id)]


# ── the rule ────────────────────────────────────────────────────────────────

def test_inheritance_is_one_more_edge_marked_inherited():
    g = _graph()
    edges = [e for e in g.edges if e.edge_type == "has_property"]
    assert len(edges) == 2
    marked = [e for e in edges if e.attributes.get("inherited")]
    assert [e.edge_source for e in marked] == ["USV1"]
    assert len([n for n in g.nodes if isinstance(n, PropertyNode)]) == 1
    assert _owners(g) == [("US5", True), ("USV1", False)]


def test_inherit_is_idempotent_and_needs_an_owner():
    g = _graph()
    inherit_property(g, "USV1", "p1")
    assert len([e for e in g.edges if e.edge_type == "has_property"]) == 2
    g.add_node(PropertyNode("p2", name="colour", value="red"))
    with pytest.raises(ValueError):
        inherit_property(g, "USV1", "p2")


def test_original_from_group_when_nothing_is_marked():
    g = Graph(graph_id="g")
    g.add_node(StratigraphicUnit("A", name="A"))
    g.add_node(StratigraphicUnit("B", name="B"))
    g.add_node(PropertyNode("p", name="material", value="x"))
    g.add_node(ParadataNodeGroup("pdg", name="A_PD"))
    g.add_edge("h1", "B", "p", "has_property")
    g.add_edge("h2", "A", "p", "has_property")
    g.add_edge("m", "p", "pdg", "is_in_paradata_nodegroup")
    g.add_edge("o", "A", "pdg", "has_paradata_nodegroup")
    assert original_owner(g, "p") == "A"


def test_original_from_creation_stamp_then_not_stated():
    g = Graph(graph_id="g")
    for u in ("A", "B"):
        g.add_node(StratigraphicUnit(u, name=u))
    g.add_node(PropertyNode("p", name="material", value="x"))
    a = g.add_edge("h1", "A", "p", "has_property")
    b = g.add_edge("h2", "B", "p", "has_property")
    assert original_owner(g, "p") is None          # edge order is not a rule
    a.attributes["created_at"] = "2026-10-04T10:00:00Z"
    b.attributes["created_at"] = "2026-10-04T09:00:00Z"
    assert original_owner(g, "p") == "B"


def test_shared_properties_lists_them():
    rows = shared_properties(_graph())
    assert rows == [{"property": "p1", "owners": ["US5", "USV1"],
                     "original": "US5"}]
    assert api.shared_properties(_graph()) == rows


def test_validate_raises_no_issue_for_an_inherited_property():
    assert api.validate(_graph())["issues"] == []


# ── the round trips ─────────────────────────────────────────────────────────

def test_emjson_roundtrip_keeps_every_owner():
    g2, _ = api.load_emjson(api.graph_to_emjson(_graph()))
    assert _owners(g2) == [("US5", True), ("USV1", False)]


def test_xlsx_roundtrip_keeps_every_owner_and_one_property(tmp_path):
    pytest.importorskip("openpyxl")
    from s3dgraphy.exporter.unified_xlsx_exporter import UnifiedXLSXExporter
    from s3dgraphy.importer.unified_xlsx_importer import UnifiedXLSXImporter
    import openpyxl

    path = tmp_path / "owners.xlsx"
    UnifiedXLSXExporter(_graph()).write(str(path))
    rows = [r for r in openpyxl.load_workbook(path)["Claims"].iter_rows(
        values_only=True)][1:]
    by_target = {r[0]: r for r in rows}
    assert by_target["US5"][1] in (None, "")       # the original: a claim
    assert by_target["USV1"][1] == "US5"           # the heir names the original
    assert by_target["USV1"][3] == "tufo"          # with a copy of the value

    g2 = UnifiedXLSXImporter(str(path), graph_id="owners").parse()
    assert _owners(g2) == [("US5", True), ("USV1", False)]


def test_xlsx_inheritance_row_without_its_original_is_kept_as_a_claim(tmp_path):
    pytest.importorskip("openpyxl")
    from s3dgraphy.importer.unified_xlsx_importer import UnifiedXLSXImporter
    import openpyxl

    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    sheets = {
        "Units": [("ID", "TYPE", "NAME"), ("US5", "US", "US5"),
                  ("USV1", "USVs", "USV1")],
        "Epochs": [("ID", "NAME", "START", "END", "COLOR")],
        "Authors": [("ID", "KIND", "DISPLAY_NAME", "ORCID", "AFFILIATION")],
        "Documents": [("ID", "FILENAME", "TITLE", "YEAR", "AUTHOR_IDS")],
        "Claims": [("TARGET_ID", "TARGET2_ID", "PROPERTY_TYPE", "VALUE"),
                   ("USV1", "US5", "material", "tufo")],
    }
    for name, rows in sheets.items():
        ws = wb.create_sheet(name)
        for r in rows:
            ws.append(r)
    path = tmp_path / "orphan.xlsx"
    wb.save(path)
    imp = UnifiedXLSXImporter(str(path), graph_id="o")
    g = imp.parse()
    assert _owners(g) == [("USV1", True)]
    assert any("inherits 'material' from 'US5'" in w for w in imp.warnings)


def test_rdf_roundtrip_keeps_every_owner_and_the_declaration(tmp_path):
    pytest.importorskip("rdflib")
    from s3dgraphy.exporter.rdf_exporter import RDFExporter
    from s3dgraphy.importer.rdf_importer import RDFImporter

    ttl = RDFExporter(str(tmp_path / "o.ttl"),
                      format="turtle").export_single_graph(_graph())
    assert "inheritsQualia" in open(ttl, encoding="utf-8").read()
    g2 = RDFImporter().parse(ttl)[0]
    assert _owners(g2) == [("US5", True), ("USV1", False)]
    # the declaration is not read back as an edge of its own
    assert len([e for e in g2.edges if e.edge_type == "has_property"]) == 2


# ── the merge does not take a shared property away ──────────────────────────

def test_merge_replacing_one_owners_qualia_keeps_the_other_owner():
    from s3dgraphy.merge.graph_merger import GraphMerger
    g = _graph()
    GraphMerger()._remove_qualia_subtree(g, g.find_node_by_id("US5"), "material")
    assert g.find_node_by_id("p1") is not None
    assert g.find_node_by_id("ext1") is not None
    assert [e.edge_source for e in g.edges
            if e.edge_type == "has_property"] == ["USV1"]
