"""The generic ``SOURCE_i`` column of the Claims sheet (E.D. 2026-09-29).

``SOURCE_i`` names what an extractor read — a document or a unit — and the
importer resolves the kind from the id. ``DOCUMENT_i`` stays valid as the older
name of the same column, read as ``SOURCE_i`` and never lost; the exporter
writes ``SOURCE_i``.
"""

import pytest

pytest.importorskip("openpyxl")

from s3dgraphy.exporter.unified_xlsx_exporter import UnifiedXLSXExporter
from s3dgraphy.graph import Graph
from s3dgraphy.importer.unified_xlsx_importer import UnifiedXLSXImporter
from s3dgraphy.nodes import StratigraphicUnit, PropertyNode, ExtractorNode
from s3dgraphy.nodes.author_node import AuthorNode
from s3dgraphy.nodes.document_node import DocumentNode


def _write(path, source_col, source_id, *, extra=None):
    from openpyxl import Workbook
    claims_h = ["TARGET_ID", "TARGET2_ID", "PROPERTY_TYPE", "VALUE", "UNITS",
                "COMBINER_REASONING", "EXTRACTOR_1", source_col, "AUTHOR_1",
                "AUTHOR_KIND_1"]
    claim = ["RSF1", "", "material", "tufo", "", "", "read off the wall",
             source_id, "A.01", "author"]
    if extra:
        for col, val in extra.items():
            claims_h.append(col)
            claim.append(val)
    sheets = {
        "Units": (["ID", "TYPE", "NAME"],
                  [["US5", "US", "US5"], ["RSF1", "RSF", "RSF1"]]),
        "Epochs": (["ID", "NAME", "START", "END", "COLOR"], []),
        "Claims": (claims_h, [claim]),
        "Authors": (["ID", "KIND", "DISPLAY_NAME", "ORCID", "AFFILIATION"],
                    [["A.01", "author", "Jane Roe", "", ""]]),
        "Documents": (["ID", "FILENAME", "TITLE", "YEAR", "AUTHOR_IDS"],
                      [["D.01", "report.pdf", "Report", "2024", "A.01"]]),
    }
    wb = Workbook()
    wb.remove(wb.active)
    for name, (headers, rows) in sheets.items():
        ws = wb.create_sheet(name)
        ws.append(headers)
        for r in rows:
            ws.append(r)
    wb.save(path)
    return path


def _by_hand(source: str) -> Graph:
    """What the row means, built without the importer."""
    g = Graph(graph_id="hand")
    g.add_node(StratigraphicUnit("US5", name="US5"))
    g.add_node(StratigraphicUnit("RSF1", name="RSF1"))
    g.add_node(AuthorNode("A.01", name="A.01"))
    g.add_node(DocumentNode("D.01", name="D.01"))
    g.add_node(PropertyNode("p", name="material", value="tufo",
                            property_type="material"))
    g.add_node(ExtractorNode("x", name=f"{source}.01",
                             description="read off the wall"))
    g.add_edge("e1", "RSF1", "p", "has_property")
    g.add_edge("e2", "p", "x", "has_data_provenance")
    g.add_edge("e3", "x", source, "extracted_from")
    g.add_edge("e4", "x", "A.01", "has_author")
    g.add_edge("e5", "D.01", "A.01", "has_author")          # Documents.AUTHOR_IDS
    return g


def _shape(g: Graph):
    """Identity-free reading: nodes by (kind, name), edges by names."""
    name = {n.node_id: n.name for n in g.nodes}
    kinds = {n.node_id: getattr(n, "node_type", type(n).__name__) for n in g.nodes}
    keep = {"has_property", "has_data_provenance", "extracted_from", "has_author"}
    nodes = sorted((kinds[n.node_id], n.name) for n in g.nodes
                   if kinds[n.node_id] in ("property", "extractor"))
    edges = sorted((name[e.edge_source], e.edge_type, name[e.edge_target])
                   for e in g.edges if e.edge_type in keep)
    return nodes, edges


def _read(path):
    imp = UnifiedXLSXImporter(str(path), graph_id="t")
    return imp.parse(), imp.warnings


def test_document_column_is_read_as_source(tmp_path):
    g, warnings = _read(_write(tmp_path / "a.xlsx", "DOCUMENT_1", "D.01"))
    assert _shape(g) == _shape(_by_hand("D.01"))
    # an old sheet is not a drifted one: no "consider updating"
    assert not any("DOCUMENT_1" in w for w in warnings)


def test_source_column_to_a_document(tmp_path):
    g, _ = _read(_write(tmp_path / "b.xlsx", "SOURCE_1", "D.01"))
    assert _shape(g) == _shape(_by_hand("D.01"))


def test_source_column_to_a_unit_names_the_extractor_by_the_unit(tmp_path):
    g, _ = _read(_write(tmp_path / "c.xlsx", "SOURCE_1", "US5"))
    assert _shape(g) == _shape(_by_hand("US5"))
    ext = next(n for n in g.nodes if getattr(n, "node_type", "") == "extractor")
    assert ext.name == "US5.01"


def test_old_and_new_give_the_same_graph(tmp_path):
    a, _ = _read(_write(tmp_path / "d1.xlsx", "DOCUMENT_1", "US5"))
    b, _ = _read(_write(tmp_path / "d2.xlsx", "SOURCE_1", "US5"))
    assert _shape(a) == _shape(b)


def test_an_unknown_source_is_reported_like_a_missing_document(tmp_path):
    g, warnings = _read(_write(tmp_path / "e.xlsx", "SOURCE_1", "AR.07"))
    assert any("unknown SOURCE 'AR.07'" in w for w in warnings)
    assert not [e for e in g.edges if e.edge_type == "extracted_from"]
    # the claim and its extractor are still there
    assert any(getattr(n, "node_type", "") == "extractor" for n in g.nodes)


def test_both_columns_in_one_sheet_lose_nothing(tmp_path):
    # SOURCE_1 empty, DOCUMENT_1 filled: the old value fills the blank
    g, _ = _read(_write(tmp_path / "f.xlsx", "SOURCE_1", "",
                        extra={"DOCUMENT_1": "D.01"}))
    assert _shape(g) == _shape(_by_hand("D.01"))
    # both filled and different: SOURCE_1 wins, said aloud
    g, warnings = _read(_write(tmp_path / "g.xlsx", "SOURCE_1", "US5",
                               extra={"DOCUMENT_1": "D.01"}))
    assert _shape(g) == _shape(_by_hand("US5"))
    assert any("disagree" in w for w in warnings)


def test_exporter_writes_source(tmp_path):
    from openpyxl import load_workbook
    path = tmp_path / "out.xlsx"
    UnifiedXLSXExporter(_by_hand("US5")).write(str(path))
    headers = [c.value for c in load_workbook(path)["Claims"][1]]
    assert "SOURCE_1" in headers and "SOURCE_2" in headers
    assert "DOCUMENT_1" not in headers
    g, _ = _read(path)
    assert _shape(g) == _shape(_by_hand("US5"))


def test_the_template_carries_source():
    from importlib.resources import files
    from openpyxl import load_workbook
    tpl = files("s3dgraphy") / "templates" / "em_data_template.xlsx"
    headers = [c.value for c in load_workbook(str(tpl))["Claims"][1]]
    assert headers == list(UnifiedXLSXImporter._COLUMNS["Claims"])
